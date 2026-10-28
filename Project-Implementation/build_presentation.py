#!/usr/bin/env python3
"""
CNAD Project Review Presentation Generator
Creates a professional 16:9 academic presentation reflecting the verified
implementation status of the Cloud-Native AI Research Discovery Platform.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Color Palette
COLOR_HEADER_BG = RGBColor(15, 23, 42)      # Slate 900
COLOR_BANNER_RED = RGBColor(168, 13, 21)     # KL University Crimson
COLOR_PRIMARY_BLUE = RGBColor(37, 99, 235)   # Royal Blue
COLOR_CARD_BG = RGBColor(248, 250, 252)      # Slate 50
COLOR_CARD_BORDER = RGBColor(226, 232, 240)  # Slate 200
COLOR_TEXT_DARK = RGBColor(15, 23, 42)       # Slate 900
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)   # Slate 500
COLOR_TEXT_WHITE = RGBColor(255, 255, 255)   # White
COLOR_SUCCESS_GREEN = RGBColor(16, 185, 129) # Emerald Green
COLOR_WARNING_AMBER = RGBColor(217, 119, 6)  # Amber
COLOR_ROW_ALT = RGBColor(241, 245, 249)      # Slate 100

LOGO_PATH = "/mnt/c/Users/Dell/Desktop/Projects/Cloud-Native AI Research Discovery Platform/ppt_assets/r1_slide1_0.png"
OUTPUT_PATH = "/mnt/c/Users/Dell/Desktop/Projects/Cloud-Native AI Research Discovery Platform/CNAD_Project_Review_Current_Status_2026-10-28.pptx"

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank layout

    def add_header(slide, title_text, category_text="CNAD Project Review — Implementation & Verification"):
        # Dark top bar
        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = COLOR_HEADER_BG
        header_shape.line.fill.background()

        # Accent red stripe
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = COLOR_BANNER_RED
        stripe.line.fill.background()

        # Category text
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(9.5), Inches(0.3))
        p_cat = tb_cat.text_frame.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_BANNER_RED
        p_cat.font.name = "Arial"

        # Main slide title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(9.5), Inches(0.65))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_WHITE
        p_title.font.name = "Arial"

        # University Logo in top right
        if os.path.exists(LOGO_PATH):
            slide.shapes.add_picture(LOGO_PATH, Inches(10.8), Inches(0.18), width=Inches(2.0))

    def add_footer(slide, slide_num):
        # Footer line
        f_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.02)
        )
        f_line.fill.solid()
        f_line.fill.fore_color.rgb = COLOR_CARD_BORDER
        f_line.line.fill.background()

        # Left footer text
        tb_left = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(8.0), Inches(0.3))
        p_l = tb_left.text_frame.paragraphs[0]
        p_l.text = "Cloud-Native AI Research Discovery Platform | Department of CSE, KL University"
        p_l.font.size = Pt(9.5)
        p_l.font.color.rgb = COLOR_TEXT_MUTED
        p_l.font.name = "Arial"

        # Right footer (Date & Slide Number)
        tb_right = slide.shapes.add_textbox(Inches(9.5), Inches(7.1), Inches(3.0), Inches(0.3))
        p_r = tb_right.text_frame.paragraphs[0]
        p_r.text = f"October 28, 2026 | Slide {slide_num}"
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.font.size = Pt(9.5)
        p_r.font.bold = True
        p_r.font.color.rgb = COLOR_TEXT_MUTED
        p_r.font.name = "Arial"

    def add_card(slide, left, top, width, height, title="", title_color=COLOR_PRIMARY_BLUE, bg_color=COLOR_CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1)

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), Inches(0.4))
            p = tb.text_frame.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = title_color
            p.font.name = "Arial"
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    # Background
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(248, 250, 252)
    bg1.line.fill.background()

    # Red Top Banner
    t_banner = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.1))
    t_banner.fill.solid()
    t_banner.fill.fore_color.rgb = COLOR_BANNER_RED
    t_banner.line.fill.background()
    tb_b = s1.shapes.add_textbox(Inches(0.9), Inches(0.65), Inches(11.533), Inches(0.8))
    p_b = tb_b.text_frame.paragraphs[0]
    p_b.text = "CLOUD-NATIVE AI RESEARCH DISCOVERY PLATFORM"
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.size = Pt(26)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_TEXT_WHITE
    p_b.font.name = "Arial"

    # Logo in header area
    if os.path.exists(LOGO_PATH):
        s1.shapes.add_picture(LOGO_PATH, Inches(10.2), Inches(0.08), width=Inches(2.2))

    # Center Review Title
    tb_rev = s1.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.333), Inches(1.2))
    p_r1 = tb_rev.text_frame.paragraphs[0]
    p_r1.text = "CNAD Project Review — Current Implementation & Demonstration"
    p_r1.alignment = PP_ALIGN.CENTER
    p_r1.font.size = Pt(24)
    p_r1.font.bold = True
    p_r1.font.color.rgb = COLOR_HEADER_BG
    p_r1.font.name = "Arial"

    p_r2 = tb_rev.text_frame.add_paragraph()
    p_r2.text = "A Modular Microservices Architecture for Academic Discovery, Automated PDF/Abstract Extraction & NLP Synthesis"
    p_r2.alignment = PP_ALIGN.CENTER
    p_r2.font.size = Pt(13)
    p_r2.font.color.rgb = COLOR_TEXT_MUTED
    p_r2.font.name = "Arial"

    p_r3 = tb_rev.text_frame.add_paragraph()
    p_r3.text = "Wednesday, October 28th, 2026"
    p_r3.alignment = PP_ALIGN.CENTER
    p_r3.font.size = Pt(14)
    p_r3.font.bold = True
    p_r3.font.color.rgb = COLOR_BANNER_RED
    p_r3.font.name = "Arial"

    # Guide Card (Left)
    add_card(s1, Inches(1.2), Inches(3.6), Inches(5.2), Inches(2.7), title="Project Guidance", title_color=COLOR_BANNER_RED)
    tb_g = s1.shapes.add_textbox(Inches(1.4), Inches(4.2), Inches(4.8), Inches(2.0))
    tf_g = tb_g.text_frame
    tf_g.word_wrap = True
    p_g1 = tf_g.paragraphs[0]
    p_g1.text = "Dr. Patha Srilatha"
    p_g1.font.size = Pt(16)
    p_g1.font.bold = True
    p_g1.font.color.rgb = COLOR_TEXT_DARK
    
    p_g2 = tf_g.add_paragraph()
    p_g2.text = "Associate Professor\nDepartment of Computer Science & Engineering\nKLH CSE Bowrampet Campus\nKL Deemed to be University, Hyderabad"
    p_g2.font.size = Pt(12)
    p_g2.font.color.rgb = COLOR_TEXT_MUTED

    # Team Card (Right)
    add_card(s1, Inches(6.9), Inches(3.6), Inches(5.2), Inches(2.7), title="Student Team Members", title_color=COLOR_PRIMARY_BLUE)
    tb_tm = s1.shapes.add_textbox(Inches(7.1), Inches(4.2), Inches(4.8), Inches(2.0))
    tf_tm = tb_tm.text_frame
    tf_tm.word_wrap = True
    members = [
        ("B. Akhil", "2420030141"),
        ("P. Ruthvik", "2420030546"),
        ("S. Akshay", "2420030579"),
    ]
    for i, (name, reg) in enumerate(members):
        p_m = tf_tm.paragraphs[0] if i == 0 else tf_tm.add_paragraph()
        p_m.text = f"•  {name}  —  {reg}"
        p_m.font.size = Pt(13)
        p_m.font.bold = True
        p_m.font.color.rgb = COLOR_TEXT_DARK
        p_m_sub = tf_tm.add_paragraph()
        p_m_sub.text = "   B.Tech Cloud-Native Application Development (CNAD)"
        p_m_sub.font.size = Pt(10.5)
        p_m_sub.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 2: PROJECT OVERVIEW & PROBLEM STATEMENT
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Project Overview & Problem Statement")
    add_footer(s2, 2)

    # 4 Challenge Cards
    cards_data = [
        ("1. Disconnected Literature Sources", 
         "Researchers waste substantial time navigating fragmented academic portals (arXiv, IEEE, CrossRef, PubMed) with disjointed search syntaxes, non-standardized schemas, and proprietary access barriers.",
         COLOR_BANNER_RED),
        ("2. Research Paper Reading Fatigue", 
         "Reading full 10-25 page papers to understand core contributions, methodologies, and findings is extremely labor-intensive. Manual extraction of key insights introduces subjective human bias.",
         COLOR_BANNER_RED),
        ("3. Contextual Discovery Deficit", 
         "Keyword search fails to match ongoing research drafts or abstracts with existing relevant literature. Researchers miss seminal papers due to terminology divergence across domains.",
         COLOR_BANNER_RED),
        ("4. Fragmented Research Management", 
         "Bookmarks, collections, and search history are scattered across browser tabs, text notes, and individual storage folders, lacking centralized cloud-native synchronization and retrieval.",
         COLOR_BANNER_RED),
    ]
    lefts = [Inches(0.8), Inches(6.8), Inches(0.8), Inches(6.8)]
    tops = [Inches(1.5), Inches(1.5), Inches(3.6), Inches(3.6)]

    for idx, (title, desc, color) in enumerate(cards_data):
        add_card(s2, lefts[idx], tops[idx], Inches(5.7), Inches(1.9), title=title, title_color=color)
        tb = s2.shapes.add_textbox(lefts[idx] + Inches(0.2), tops[idx] + Inches(0.55), Inches(5.3), Inches(1.25))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Integrated Solution Banner
    sol_card = add_card(s2, Inches(0.8), Inches(5.7), Inches(11.733), Inches(1.15), title="", bg_color=RGBColor(238, 242, 255))
    tb_sol = s2.shapes.add_textbox(Inches(1.0), Inches(5.75), Inches(11.333), Inches(1.0))
    tf_sol = tb_sol.text_frame
    tf_sol.word_wrap = True
    p_s1 = tf_sol.paragraphs[0]
    p_s1.text = "THE PROPOSED PLATFORM SOLUTION:"
    p_s1.font.size = Pt(11.5)
    p_s1.font.bold = True
    p_s1.font.color.rgb = COLOR_PRIMARY_BLUE
    p_s2 = tf_sol.add_paragraph()
    p_s2.text = "A unified cloud-native environment integrating live academic search (OpenAlex), dual-mode document intake (raw PDF and pasted abstract), structured scientific NLP extraction, contextual Jaccard recommendations, and persistent library storage in a decoupled microservices architecture."
    p_s2.font.size = Pt(11)
    p_s2.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 3: OBJECTIVES & RESEARCH GAP
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Objectives & Problem Gap Identification")
    add_footer(s3, 3)

    # Left: Objectives
    add_card(s3, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Platform Objectives (Implemented)", title_color=COLOR_PRIMARY_BLUE)
    tb_obj = s3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_obj = tb_obj.text_frame
    tf_obj.word_wrap = True
    objectives = [
        ("Centralized Academic Discovery", "Provide a single interface querying 250M+ open scholarly papers via OpenAlex with automatic local catalog fallback."),
        ("Dual Ingestion Pipeline", "Accept research abstracts and full text-based PDF papers with real-time progress tracking."),
        ("Automated Scientific Section Extraction", "Extract structured research problems, objectives, methodology, quantitative findings, limitations, gaps, and keywords."),
        ("Contextual Similarity Recommendations", "Compute literature recommendations from working abstracts using transparent Jaccard term overlap."),
        ("Binary Document Streaming", "Enable seamless PDF storage and in-browser streaming via S3-compatible MinIO object storage."),
        ("Cloud-Native Microservices Architecture", "Decouple services into containerized, independently scalable units with health checks and auto-restart policies."),
    ]
    for i, (head, body) in enumerate(objectives):
        p = tf_obj.paragraphs[0] if i == 0 else tf_obj.add_paragraph()
        p.text = f"• {head}: "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT_DARK
        p.font.name = "Arial"
        # Add normal text
        run = p.add_run()
        run.text = body
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Research Gap Evolution Table
    add_card(s3, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Literature Survey & Identified Research Gap", title_color=COLOR_BANNER_RED)
    tb_lit = s3.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_lit = tb_lit.text_frame
    tf_lit.word_wrap = True
    lit_items = [
        ("Kratzke & Quint (2017) [1]", "Identified microservices and containers as essential for cloud scalability, but noted gaps in decoupled scholarly tool deployment."),
        ("Di Francesco et al. (2019) [2]", "Demonstrated that service decomposition improves reliability, yet highlighted challenges in service orchestration and asynchronous workflows."),
        ("Asyrofi et al. (2021) [3]", "Surveyed paper recommendation systems, emphasizing the need for transparent, explainable recommendations over black-box proprietary ranking."),
        ("The Concrete Problem Gap", "Existing tools separate search (Google Scholar), reference management (Zotero), and document analysis. No unified, containerized platform combines real-time discovery, structural NLP extraction, object storage, and contextual recommendations in an open cloud-native microservices architecture."),
    ]
    for i, (head, body) in enumerate(lit_items):
        p = tf_lit.paragraphs[0] if i == 0 else tf_lit.add_paragraph()
        p.text = f"• {head}: "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_BANNER_RED if i == 3 else COLOR_TEXT_DARK
        run = p.add_run()
        run.text = body
        run.font.bold = (i == 3)
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT_DARK if i == 3 else COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "System Architecture & Microservices Request Flow")
    add_footer(s4, 4)

    # Diagram Cards
    # Top Tier: Frontend
    add_card(s4, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.9), title="Client Layer: React 18 + Vite SPA (:5173)", title_color=COLOR_PRIMARY_BLUE)
    tb_fe = s4.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(11.3), Inches(0.45))
    tb_fe.text_frame.word_wrap = True
    p_fe = tb_fe.text_frame.paragraphs[0]
    p_fe.text = "Modules: Dashboard | Academic Search (OpenAlex) | Dual-Mode Upload (PDF & Abstract) | Paper Details & AI View | Library Bookmarks"
    p_fe.font.size = Pt(10.5)
    p_fe.font.color.rgb = COLOR_TEXT_DARK

    # Mid Tier: Gateway
    add_card(s4, Inches(0.8), Inches(2.55), Inches(11.733), Inches(0.85), title="API Gateway Reverse Proxy Layer: Express Gateway (:3000)", title_color=COLOR_HEADER_BG)
    tb_gw = s4.shapes.add_textbox(Inches(1.0), Inches(2.95), Inches(11.3), Inches(0.4))
    tb_gw.text_frame.word_wrap = True
    p_gw = tb_gw.text_frame.paragraphs[0]
    p_gw.text = "Features: Single Entrypoint | JWT Header Forwarding | Cluster Health Check (/health) | Helmet Cross-Origin Resource Policy (PDF Stream)"
    p_gw.font.size = Pt(10.5)
    p_gw.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Tier: 5 Core Services (Grid of 5)
    srvs = [
        ("Auth Service\n(:3001, Node.js)", "• BCrypt (10 rounds)\n• Stateless JWT (7d)\n• RBAC Middleware\n• User schema validation", COLOR_BANNER_RED),
        ("Paper Service\n(:3002, Node.js)", "• Paper CRUD & queries\n• Abstract ingestion\n• MinIO S3 SDK client\n• Re-analysis coordinator", COLOR_PRIMARY_BLUE),
        ("Search & Rec Service\n(:3003, Node.js)", "• OpenAlex REST API\n• Local DB Fallback\n• Jaccard term matching\n• Search history logs", COLOR_PRIMARY_BLUE),
        ("PDF Processing Service\n(:8001, FastAPI)", "• PyMuPDF text stream\n• S3 byte downloading\n• Malformed PDF trapping\n• Zero-text detection", RGBColor(16, 185, 129)),
        ("AI Analysis Service\n(:8002, FastAPI)", "• Scientific NLP engine\n• 7-section extraction\n• LLM fallback gateway\n• Async webhook callback", RGBColor(16, 185, 129)),
    ]
    w_srv = Inches(2.22)
    for idx, (title, body, color) in enumerate(srvs):
        l = Inches(0.8) + idx * (w_srv + Inches(0.15))
        add_card(s4, l, Inches(3.6), w_srv, Inches(1.9), title=title, title_color=color)
        tb_s = s4.shapes.add_textbox(l + Inches(0.1), Inches(4.2), w_srv - Inches(0.2), Inches(1.2))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        p_s = tf_s.paragraphs[0]
        p_s.text = body
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Storage Tier
    add_card(s4, Inches(0.8), Inches(5.65), Inches(5.7), Inches(1.25), title="Persistence: MongoDB Atlas / Engine (:27017)", title_color=RGBColor(5, 150, 105))
    tb_db = s4.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(5.3), Inches(0.75))
    tb_db.text_frame.word_wrap = True
    p_db = tb_db.text_frame.paragraphs[0]
    p_db.text = "Collections: users, papers, paperanalyses, bookmarks, recommendations, searchhistories.\nIndexes: Text search ($text on title, abstract, keywords), sparse DOIs, user bookmarks."
    p_db.font.size = Pt(9.5)
    p_db.font.color.rgb = COLOR_TEXT_DARK

    add_card(s4, Inches(6.8), Inches(5.65), Inches(5.7), Inches(1.25), title="Object Storage: MinIO S3 API (:9000 / :9001)", title_color=COLOR_BANNER_RED)
    tb_s3 = s4.shapes.add_textbox(Inches(7.0), Inches(6.05), Inches(5.3), Inches(0.75))
    tb_s3.text_frame.word_wrap = True
    p_s3 = tb_s3.text_frame.paragraphs[0]
    p_s3.text = "Bucket: research-papers | Cloud-Native S3-compatible binary blob storage.\nAccess: Automated initialization (minio-setup), presigned URLs, inline stream proxying."
    p_s3.font.size = Pt(9.5)
    p_s3.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 5: TECHNOLOGY STACK & ARCHITECTURAL RATIONALE
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Technology Stack & Architectural Rationale")
    add_footer(s5, 5)

    # Table of Tech Stack
    rows = 7
    cols = 4
    left = Inches(0.8)
    top = Inches(1.5)
    width = Inches(11.733)
    height = Inches(5.3)

    table_shape = s5.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(2.5)
    table.columns[2].width = Inches(4.8)
    table.columns[3].width = Inches(2.233)

    headers = ["Component Tier", "Selected Technology", "Architectural Rationale & Purpose", "Verification Status"]
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = h_text
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER_BG
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_WHITE
        p.alignment = PP_ALIGN.CENTER

    data = [
        ("Frontend UI", "React 18 + Vite", "Component-driven SPA, fast HMR, 925ms production build, responsive dual-tab stepper.", "Implemented & Verified"),
        ("API Gateway", "Express + Proxy", "Single entrypoint, stateless token verification, Helmet CORP for PDF streams.", "Implemented & Verified"),
        ("Backend Services", "Node.js (v20) + Express", "Asynchronous non-blocking I/O ideal for metadata handling and REST endpoints.", "Implemented & Verified"),
        ("PDF & AI Services", "Python 3.10 + FastAPI", "High throughput async engine with native C-bindings for PyMuPDF and NLP computation.", "Implemented & Verified"),
        ("Document DB", "MongoDB (Mongoose)", "Flexible schema accommodating nested authors, analytical fields, and text indexes.", "Implemented & Verified"),
        ("Object Storage", "MinIO (S3 API)", "Cloud-native S3-compatible binary PDF storage decoupling files from compute containers.", "Implemented & Verified"),
    ]

    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, cell_data in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = cell_data
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_ROW_ALT if row_idx % 2 == 1 else COLOR_TEXT_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_SUCCESS_GREEN if col_idx == 3 else COLOR_TEXT_DARK
            p.font.bold = (col_idx == 1 or col_idx == 3)
            if col_idx == 3:
                p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 6: RESEARCH DISCOVERY WORKFLOW
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Academic Research Discovery Workflow")
    add_footer(s6, 6)

    # 3 Column Workflow
    w_col = Inches(3.7)
    gap = Inches(0.3)

    # Step 1: OpenAlex Integration
    add_card(s6, Inches(0.8), Inches(1.5), w_col, Inches(4.0), title="1. OpenAlex Live Query Engine", title_color=COLOR_PRIMARY_BLUE)
    tb_s1 = s6.shapes.add_textbox(Inches(1.0), Inches(2.1), w_col - Inches(0.4), Inches(3.2))
    tf_s1 = tb_s1.text_frame
    tf_s1.word_wrap = True
    p = tf_s1.paragraphs[0]
    p.text = "• Live Academic Search API:\nConnects to OpenAlex REST API indexing 250M+ scholarly works across all scientific domains.\n\n• Metadata Extraction:\nParses work title, authors, affiliations, publication year, publication venue/journal, open-access PDF links, and DOIs.\n\n• Paginated Retrieval:\nSupports configurable page numbers and limits up to 50 results per query with relevance sorting."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    # Step 2: Resilient Local Fallback
    add_card(s6, Inches(0.8) + w_col + gap, Inches(1.5), w_col, Inches(4.0), title="2. Resilient Catalog Fallback", title_color=COLOR_BANNER_RED)
    tb_s2 = s6.shapes.add_textbox(Inches(1.0) + w_col + gap, Inches(2.1), w_col - Inches(0.4), Inches(3.2))
    tf_s2 = tb_s2.text_frame
    tf_s2.word_wrap = True
    p = tf_s2.paragraphs[0]
    p.text = "• Automatic Fault Detection:\nIf OpenAlex times out, reaches rate limits, or is offline, the gateway catches the error instantly.\n\n• Local MongoDB Fallback:\nExecutes multi-field regex and text index queries across local papers (title, abstract, keywords, author names).\n\n• Zero Downtime:\nReturns local catalog papers with an informational notice flag to ensure uninterrupted research workflows."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    # Step 3: Discovery Presentation & Actions
    add_card(s6, Inches(0.8) + (w_col + gap) * 2, Inches(1.5), w_col, Inches(4.0), title="3. Interactive Exploration", title_color=RGBColor(16, 185, 129))
    tb_s3 = s6.shapes.add_textbox(Inches(1.0) + (w_col + gap) * 2, Inches(2.1), w_col - Inches(0.4), Inches(3.2))
    tf_s3 = tb_s3.text_frame
    tf_s3.word_wrap = True
    p = tf_s3.paragraphs[0]
    p.text = "• Verified Results Verification:\nVerified against queries such as 'cloud native microservices' returning 10+ live papers with complete citation metadata.\n\n• Seamless Library Ingestion:\nUsers can click 'Save to Library' to import external research records directly into MongoDB.\n\n• Contextual Bridging:\nImported papers become available for contextual recommendation matching and local analysis."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    # Summary Callout Banner
    add_card(s6, Inches(0.8), Inches(5.7), Inches(11.733), Inches(1.15), title="", bg_color=RGBColor(240, 253, 244))
    tb_sum = s6.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.333), Inches(0.9))
    tf_sum = tb_sum.text_frame
    tf_sum.word_wrap = True
    p = tf_sum.paragraphs[0]
    p.text = "VERIFIED IN E2E SUITE (STEP 4 & 15):"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(5, 150, 105)
    p2 = tf_sum.add_paragraph()
    p2.text = "OpenAlex query returned 10 real papers (Total indexed: 19,109 matches for microservices). When switching to local catalog, multi-field regex returned verified uploaded records with 100% precision."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 7: ABSTRACT SUBMISSION & RESEARCH ANALYSIS
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Abstract Submission & Scientific AI Analysis Pipeline", "Highest Priority Demonstration Feature")
    add_footer(s7, 7)

    # Left: Workflow & Stepper
    add_card(s7, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Dual-Mode Ingest & Asynchronous Processing", title_color=COLOR_BANNER_RED)
    tb_a1 = s7.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(5.3), Inches(4.7))
    tf_a1 = tb_a1.text_frame
    tf_a1.word_wrap = True
    p = tf_a1.paragraphs[0]
    p.text = "1. Submission & Input Validation (Frontend & Gateway):\n• Users paste raw abstract text into the specialized Upload tab.\n• Input validation verifies non-empty text; blank input returns HTTP 400.\n• Dynamic title extraction parses first sentence if omitted.\n\n2. Asynchronous Dispatch & Status Progression:\n• Creates MongoDB record with source='abstract' and status='analyzing'.\n• Paper Service triggers AI Service via internal network POST.\n• UI displays 4-stage stepper: Validation → Ingest → Section Analysis → Synthesis.\n\n3. Local Scientific NLP Engine (nlp_analyzer.py):\n• Executed automatically when no external API key is set.\n• Deterministic extractive pipeline analyzing scientific sentence syntax.\n• Explicitly credited as 'Local Scientific NLP Pipeline' (no fake LLM claims)."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    # Right: The 7 Extracted Structured Sections
    add_card(s7, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Generated Structured Scientific Analysis", title_color=COLOR_PRIMARY_BLUE)
    tb_a2 = s7.shapes.add_textbox(Inches(7.0), Inches(1.95), Inches(5.3), Inches(4.7))
    tf_a2 = tb_a2.text_frame
    tf_a2.word_wrap = True
    sections = [
        ("Research Problem", "Identifies core challenges, bottlenecks, or constraints addressed in the work."),
        ("Primary Objectives", "Synthesizes primary research aims, architectural goals, and hypotheses."),
        ("Methodology", "Extracts experimental setups, algorithms, frameworks, and benchmark designs."),
        ("Major Findings", "Scores and extracts quantitative achievements (e.g., '34% tail latency reduction, 21% cost drop')."),
        ("Limitations & Gaps", "Identifies explicitly stated constraints, threats to validity, and scoped edge cases."),
        ("Future Directions", "Extracts author recommendations for forward-looking investigations."),
        ("Keywords & Domain", "Frequency-weighted terminology vector and domain classification (Cloud Computing, AI)."),
    ]
    for i, (head, desc) in enumerate(sections):
        p = tf_a2.paragraphs[0] if i == 0 else tf_a2.add_paragraph()
        p.text = f"✔ {head}: "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_PRIMARY_BLUE
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 8: PDF UPLOAD & PROCESSING WORKFLOW
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "PDF Upload, MinIO Storage & Text Extraction Pipeline")
    add_footer(s8, 8)

    # 4 Steps Cards Across
    w4 = Inches(2.75)
    gap4 = Inches(0.24)
    steps_data = [
        ("1. Multipart Upload", "• 50MB file size limit\n• MIME type validation (application/pdf)\n• Uploaded to MinIO bucket via AWS SDK client\n• Unique key: papers/<uid>/<ts>-<name>.pdf", COLOR_PRIMARY_BLUE),
        ("2. Text Extraction", "• Fast C-native PyMuPDF (fitz) extraction\n• Handles complex multi-column typography\n• Preserves section flows and technical tables\n• Asynchronous webhook updates", COLOR_PRIMARY_BLUE),
        ("3. Edge Case Handling", "• Malformed PDFs caught safely (status: 'failed')\n• Scanned/image-only PDFs detected (<50 chars)\n• Informative status: 'completed_without_ai'\n• Guides user to submit OCR-processed text", COLOR_BANNER_RED),
        ("4. S3 Binary Stream", "• Endpoint: GET /api/papers/:id/pdf\n• Streams binary from MinIO through Gateway\n• Helmet crossOriginResourcePolicy header\n• In-browser inspection without external tools", RGBColor(16, 185, 129)),
    ]

    for idx, (title, body, color) in enumerate(steps_data):
        l = Inches(0.8) + idx * (w4 + gap4)
        add_card(s8, l, Inches(1.5), w4, Inches(3.8), title=title, title_color=color)
        tb = s8.shapes.add_textbox(l + Inches(0.1), Inches(2.1), w4 - Inches(0.2), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = body
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Banner
    add_card(s8, Inches(0.8), Inches(5.55), Inches(11.733), Inches(1.3), title="Verified Execution in End-to-End Test Suite", title_color=COLOR_HEADER_BG)
    tb_b = s8.shapes.add_textbox(Inches(1.0), Inches(5.95), Inches(11.333), Inches(0.8))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "• Valid PDF Upload (Step 9-11): Stored in MinIO, extracted by PyMuPDF, analyzed by AI service, streamed 1,142 bytes back to client.\n• Corrupt PDF Handling (Step 12): Malformed file trapped gracefully without process termination, marked 'failed'.\n• Scanned PDF Handling (Step 13): Zero-text PDF detected and marked 'completed_without_ai' with actionable feedback."
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 9: RELATED PAPER RECOMMENDATIONS & LIBRARY
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Contextual Recommendations & Personal Research Library")
    add_footer(s9, 9)

    # Left: Recommendation Engine
    add_card(s9, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Contextual Recommendation Engine", title_color=COLOR_PRIMARY_BLUE)
    tb_r1 = s9.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(5.3), Inches(4.7))
    tf_r1 = tb_r1.text_frame
    tf_r1.word_wrap = True
    p = tf_r1.paragraphs[0]
    p.text = "• Transparent Jaccard Similarity Model:\nUnlike opaque black-box recommendations, the platform implements an explainable set-theoretic similarity function:\n\n        Similarity(A, B) = |A ∩ B| / |A ∪ B|\n\n• Term Extraction from Working Abstracts:\nStopwords are stripped ('the', 'paper', 'approach') and domain content terms are scored to form the query keyword vector.\n\n• Candidate Ranking & Scored Matches:\nMatches local candidate literature and scores between 0.0 and 1.0 based on keyword overlap (70%) and research domain alignment (30%).\n\n• Explainable Match Reasons:\nResults explicitly state matching concepts (e.g., 'Shares keywords: Kubernetes, Microservices, Autoscaling')."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    # Right: Research Library & Persistence
    add_card(s9, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Personal Research Library Management", title_color=COLOR_BANNER_RED)
    tb_r2 = s9.shapes.add_textbox(Inches(7.0), Inches(1.95), Inches(5.3), Inches(4.7))
    tf_r2 = tb_r2.text_frame
    tf_r2.word_wrap = True
    p = tf_r2.paragraphs[0]
    p.text = "• One-Click Research Bookmarking:\nUsers save discovered or uploaded papers directly into their personal MongoDB library collection.\n\n• Multi-Session Persistence:\nBookmarked papers, AI analytical summaries, and S3 file keys persist across browser refreshes and user logouts.\n\n• Duplicate Prevention:\nCompound unique indexes on (userId, paperId) ensure no duplicate bookmarks clutter user collections.\n\n• Direct Detail Navigation:\nUsers open saved papers directly from the library to review AI-generated problem statements, findings, and original PDF streams.\n\n• Verified Status:\nTested and confirmed in Step 14 of E2E verification suite."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 10: SECURITY, PERSISTENCE & RELIABILITY
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Security, Persistence & Cloud-Native Reliability")
    add_footer(s10, 10)

    # 3 Category Cards
    w3 = Inches(3.7)
    gap3 = Inches(0.3)

    add_card(s10, Inches(0.8), Inches(1.5), w3, Inches(5.35), title="Security & Access Control", title_color=COLOR_BANNER_RED)
    tb_sec = s10.shapes.add_textbox(Inches(1.0), Inches(2.1), w3 - Inches(0.4), Inches(4.5))
    tf_sec = tb_sec.text_frame
    tf_sec.word_wrap = True
    p = tf_sec.paragraphs[0]
    p.text = "• Password Hashing:\nBCrypt with 10 salt rounds ensures no plain text credentials are stored.\n\n• Stateless JWT Tokens:\nSigned with HMAC-SHA256, verified by API Gateway and downstream services.\n\n• Role-Based Access Control:\nMiddleware enforces permissions across student, researcher, and admin roles.\n\n• Input Validation:\nStrict schema enforcement on abstracts (10,000 char buffer) and PDF files (50MB cap).\n\n• CORS & Helmet Protection:\nRestricts unauthorized origins while allowing cross-origin PDF streaming."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    add_card(s10, Inches(0.8) + w3 + gap3, Inches(1.5), w3, Inches(5.35), title="Multi-Tier Persistence", title_color=COLOR_PRIMARY_BLUE)
    tb_per = s10.shapes.add_textbox(Inches(1.0) + w3 + gap3, Inches(2.1), w3 - Inches(0.4), Inches(4.5))
    tf_per = tb_per.text_frame
    tf_per.word_wrap = True
    p = tf_per.paragraphs[0]
    p.text = "• MongoDB Storage:\nStores user profiles, academic metadata, structured AI analyses, and bookmarks.\n\n• Persistent Volumes:\nDocker named volume 'mongo-data' ensures database survives container rebuilds.\n\n• MinIO Object Storage:\nS3-compatible bucket 'research-papers' with persistent volume 'minio-data'.\n\n• Relational Referencing:\nMongoose references link Paper records to PaperAnalysis documents with compound index protection."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    add_card(s10, Inches(0.8) + (w3 + gap3) * 2, Inches(1.5), w3, Inches(5.35), title="Cloud-Native Reliability", title_color=RGBColor(16, 185, 129))
    tb_rel = s10.shapes.add_textbox(Inches(1.0) + (w3 + gap3) * 2, Inches(2.1), w3 - Inches(0.4), Inches(4.5))
    tf_rel = tb_rel.text_frame
    tf_rel.word_wrap = True
    p = tf_rel.paragraphs[0]
    p.text = "• Service Health Probes:\nAPI Gateway exposes /health aggregating live downstream probes for all 5 microservices.\n\n• Container Restart Policies:\nPersistent services configured with 'restart: unless-stopped' for auto-recovery.\n\n• Isolated Bridge Network:\nServices communicate via internal Docker DNS names ('mongodb', 'paper-service').\n\n• Non-Blocking Timeouts:\nPyMuPDF, AI analysis, and OpenAlex calls protected by circuit timeouts to prevent cascading failure."
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 11: IMPLEMENTATION & VERIFICATION RESULTS
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Implementation & Automated Verification Results")
    add_footer(s11, 11)

    # Verification Table
    t_shape = s11.shapes.add_table(8, 4, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.35))
    vt = t_shape.table
    vt.columns[0].width = Inches(2.4)
    vt.columns[1].width = Inches(2.8)
    vt.columns[2].width = Inches(4.333)
    vt.columns[3].width = Inches(2.2)

    v_headers = ["Verification Scope", "Test Suite / Metric", "Observed Test Outcome", "Result Status"]
    for col_idx, h_text in enumerate(v_headers):
        cell = vt.cell(0, col_idx)
        cell.text = h_text
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER_BG
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_WHITE
        p.alignment = PP_ALIGN.CENTER

    v_data = [
        ("Frontend Build", "Vite Production Bundler", "Zero syntax errors, 1936 modules transformed, built in 925ms.", "PASSED (0 Errors)"),
        ("Auth Service Unit Tests", "Jest + Supertest Suite", "7 / 7 tests passed (JWT verify, expiry, RBAC roles, BCrypt hashing).", "7 / 7 PASSED"),
        ("Paper Service Unit Tests", "Jest Model & Schema Suite", "11 / 11 tests passed (Abstract source, status enum, Joi validations).", "11 / 11 PASSED"),
        ("Search Service Unit Tests", "Jest Recommendation Suite", "11 / 11 tests passed (Jaccard similarity, stopword filtering, schemas).", "11 / 11 PASSED"),
        ("E2E Integration Suite", "python -u test_e2e.py", "17 / 17 steps passed through API Gateway (port 3000).", "17 / 17 PASSED"),
        ("Abstract & AI Pipeline", "End-to-End Ingest & Poll", "Extracted Problem, Objectives, Methodology, Findings, Limitations, Gaps.", "VERIFIED"),
        ("PDF Ingest & Storage", "MinIO + PyMuPDF Extraction", "Uploaded, extracted text, handled malformed PDF, streamed binary PDF.", "VERIFIED"),
    ]

    for row_idx, row in enumerate(v_data, start=1):
        for col_idx, text in enumerate(row):
            cell = vt.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_ROW_ALT if row_idx % 2 == 1 else COLOR_TEXT_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_SUCCESS_GREEN if col_idx == 3 else COLOR_TEXT_DARK
            p.font.bold = (col_idx == 0 or col_idx == 3)
            if col_idx == 3:
                p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 12: CLOUD-NATIVE IMPLEMENTATION: CURRENT STATUS
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Cloud-Native Implementation: Current Status Matrix")
    add_footer(s12, 12)

    # Matrix Table
    m_shape = s12.shapes.add_table(11, 4, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.35))
    mt = m_shape.table
    mt.columns[0].width = Inches(2.8)
    mt.columns[1].width = Inches(2.2)
    mt.columns[2].width = Inches(4.533)
    mt.columns[3].width = Inches(2.2)

    m_headers = ["Cloud-Native Dimension", "Technology Component", "Current Implementation Detail", "Implementation State"]
    for col_idx, h_text in enumerate(m_headers):
        cell = mt.cell(0, col_idx)
        cell.text = h_text
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER_BG
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_WHITE
        p.alignment = PP_ALIGN.CENTER

    m_data = [
        ("Containerization", "Docker & Docker Compose", "All 9 services containerized with custom Dockerfiles & healthchecks.", "Implemented & Verified"),
        ("Microservice Architecture", "5 Discrete Node/Python Svcs", "Decoupled services communicating via internal Docker bridge network.", "Implemented & Verified"),
        ("API Gateway Proxy", "Express Gateway (:3000)", "Central reverse proxy with route rewriting, auth forwarding, and Helmet.", "Implemented & Verified"),
        ("Object Storage", "MinIO Cluster (:9000)", "S3-compatible bucket storage for binary PDFs with presigned streaming.", "Implemented & Verified"),
        ("Document Database", "MongoDB (:27017)", "Multi-collection schema with persistent disk volume (mongo-data).", "Implemented & Verified"),
        ("Local Scientific AI", "Extractive NLP Pipeline", "Deterministic sentence scoring for research section extraction.", "Implemented & Verified"),
        ("Kubernetes Manifests", "Deployments & Services", "k8s YAML manifests structured for cluster deployment.", "Configured / Prototype"),
        ("Automated CI/CD", "GitHub Actions / NPM Scripts", "Automated unit test runner and production build pipelines configured.", "Configured / Ready"),
        ("GitOps Deployment", "Argo CD / Helm", "Automated continuous delivery to Kubernetes clusters.", "Planned (Phase 5)"),
        ("Cloud Deployment", "AWS EKS / CloudWatch", "Managed Kubernetes cluster deployment on AWS infrastructure.", "Planned (Phase 6)"),
    ]

    for row_idx, row in enumerate(m_data, start=1):
        for col_idx, text in enumerate(row):
            cell = mt.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_ROW_ALT if row_idx % 2 == 1 else COLOR_TEXT_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9)
            if col_idx == 3:
                p.alignment = PP_ALIGN.CENTER
                p.font.bold = True
                if "Implemented" in text:
                    p.font.color.rgb = COLOR_SUCCESS_GREEN
                elif "Configured" in text:
                    p.font.color.rgb = COLOR_PRIMARY_BLUE
                else:
                    p.font.color.rgb = COLOR_WARNING_AMBER
            else:
                p.font.color.rgb = COLOR_TEXT_DARK
                p.font.bold = (col_idx == 0)

    # ==========================================
    # SLIDE 13: REMAINING WORK & FUTURE ENHANCEMENTS
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Remaining Work, Limitations & Future Enhancements")
    add_footer(s13, 13)

    # Left: Current Limitations (Honest Assessment)
    add_card(s13, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Current Prototype Limitations", title_color=COLOR_BANNER_RED)
    tb_lim = s13.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_lim = tb_lim.text_frame
    tf_lim.word_wrap = True
    limits = [
        ("External LLM API Key Requirement", "Generative paper synthesis currently requires a configured external AI_API_KEY (OpenAI/Groq). In its absence, the verified Local Scientific NLP Pipeline operates as a reliable extractive fallback."),
        ("OCR for Scanned Historical Papers", "PyMuPDF extracts embedded font streams. Purely scanned image-only PDFs currently return 'completed_without_ai' and prompt the user to upload text. Dedicated OCR integration is scheduled."),
        ("Local Docker Environment", "The current demonstration operates reliably on local Docker Compose. Production deployment on a managed public cloud Kubernetes cluster (AWS EKS) remains the target for Phase 6."),
        ("Semantic Vector Search", "Recommendations currently use Jaccard keyword overlap. Upgrading to high-dimensional dense vector embeddings (e.g., Pinecone/Milvus) will further improve conceptual nuance."),
    ]
    for i, (head, body) in enumerate(limits):
        p = tf_lim.paragraphs[0] if i == 0 else tf_lim.add_paragraph()
        p.text = f"• {head}:\n"
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_BANNER_RED
        run = p.add_run()
        run.text = f"  {body}\n"
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Phase 5 & 6 Roadmap
    add_card(s13, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Phased Development Roadmap", title_color=COLOR_PRIMARY_BLUE)
    tb_rd = s13.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_rd = tb_rd.text_frame
    tf_rd.word_wrap = True
    roadmap = [
        ("Phase 5: Helm, GitOps & Security Pipeline", "• Package microservices into modular Helm charts.\n• Implement Argo CD for declarative GitOps synchronization.\n• Integrate Trivy container vulnerability scanning in GitHub Actions."),
        ("Phase 6: AWS EKS Cloud Deployment", "• Provision managed AWS EKS cluster with multi-AZ node groups.\n• Migrate local MinIO to managed Amazon S3 with IAM role security.\n• Configure AWS ALB Ingress Controller with SSL termination."),
        ("Observability & Metrics", "• Deploy Prometheus operators for cluster resource telemetry.\n• Build Grafana dashboards tracking API request latency, PyMuPDF extraction times, and container memory utilization."),
        ("Dense Vector Embeddings", "• Integrate sentence-transformer embeddings for semantic paper similarity matching beyond lexical keyword overlap."),
    ]
    for i, (head, body) in enumerate(roadmap):
        p = tf_rd.paragraphs[0] if i == 0 else tf_rd.add_paragraph()
        p.text = f"• {head}:\n"
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_PRIMARY_BLUE
        run = p.add_run()
        run.text = f"{body}\n"
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 14: CONCLUSION & LIVE DEMONSTRATION FLOW
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Conclusion & Live Demonstration Sequence")
    add_footer(s14, 14)

    # Left: Conclusion Summary
    add_card(s14, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Project Summary & Achievements", title_color=COLOR_BANNER_RED)
    tb_c1 = s14.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    conclusions = [
        ("From Concept to Working System", "Successfully transitioned from initial proposal into an active, containerized, multi-service platform."),
        ("Verified Dual Ingest Engine", "Both raw text abstracts and PDF research papers are reliably processed with zero simulated mock data."),
        ("Genuine Scientific Section Analysis", "Generates real problem statements, objectives, methodology, quantitative findings, limitations, and keywords."),
        ("Cloud-Native Architecture", "Strict microservice decoupling, stateless gateway routing, and persistent MinIO/MongoDB storage."),
        ("100% Test Passing Rate", "29 unit tests passed; 17 end-to-end integration workflows verified across all container boundaries."),
    ]
    for i, (head, body) in enumerate(conclusions):
        p = tf_c1.paragraphs[0] if i == 0 else tf_c1.add_paragraph()
        p.text = f"✔ {head}:\n"
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_BANNER_RED
        run = p.add_run()
        run.text = f"  {body}\n"
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT_DARK

    # Right: Live Demonstration Sequence
    add_card(s14, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.35), title="Live Demonstration Sequence for Review", title_color=RGBColor(16, 185, 129))
    tb_c2 = s14.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    demo_steps = [
        ("Step 1: User Login", "Log in via demo@university.edu / password123. Show JWT token authentication and dashboard metrics."),
        ("Step 2: Academic Discovery", "Search 'microservices' in Discover page. Show live OpenAlex results and local fallback capability."),
        ("Step 3: Abstract Submission", "Paste real Kubernetes research abstract. Show 4-stage processing stepper and extracted findings."),
        ("Step 4: PDF Document Upload", "Upload research paper PDF. Show MinIO S3 key creation, PyMuPDF text extraction, and inline PDF view."),
        ("Step 5: Recommendations & Library", "Show contextual paper recommendations matching the abstract. Bookmark paper and verify persistence."),
    ]
    for i, (step, desc) in enumerate(demo_steps):
        p = tf_c2.paragraphs[0] if i == 0 else tf_c2.add_paragraph()
        p.text = f"{step}:\n"
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = RGBColor(5, 150, 105)
        run = p.add_run()
        run.text = f"  {desc}\n"
        run.font.bold = False
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Verified URLs
    p_url = tf_c2.add_paragraph()
    p_url.text = "Verified Running URLs: Frontend: http://localhost:5173 | Gateway: http://localhost:3000 | Health: http://localhost:3000/health"
    p_url.font.size = Pt(9.5)
    p_url.font.bold = True
    p_url.font.color.rgb = COLOR_PRIMARY_BLUE

    # Save presentation
    prs.save(OUTPUT_PATH)
    print(f"Successfully generated 16:9 presentation: {OUTPUT_PATH}")
    print(f"Total slides: {len(prs.slides)}")

if __name__ == "__main__":
    create_deck()
